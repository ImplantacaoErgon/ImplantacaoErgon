#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Importador do cronograma real da PCR (export .xlsx do MS Project) para o Ergon PM.

O QUE ESTE SCRIPT FAZ
----------------------
Lê um arquivo .xlsx exportado do MS Project (colunas: EDT, Id, Nome da tarefa,
Início, Término, Duração, % concluída, Início/Término da Linha de Base,
Predecessoras, Nomes dos recursos, Variação do término) e recria a mesma
estrutura no Ergon PM via a API REST do backend: cria o projeto, as Etapas,
as Frentes de Trabalho, os Recursos (responsáveis), os Tipos de Atividade que
faltarem, todas as Atividades (preservando a hierarquia WBS via
atividade_pai_id), as Dependências entre elas, e extrai a única tarefa de
duração zero como Marco.

Este script é ESPECÍFICO para a estrutura desse cronograma da PCR (WBS de
4 fases no nível 1 — Planejamento/Iniciação/Execução/Encerramento — e um
conjunto fixo de Frentes no nível 2 dentro de "Execução", mapeado abaixo em
FRENTE_POR_EDT2). Para importar um cronograma com WBS diferente, ajuste esse
mapeamento antes de rodar.

DECISÕES JÁ COMBINADAS (não pedem confirmação ao rodar):
- Etapas recriadas para bater com as 4 fases reais do cronograma (não usa as
  3 etapas fictícias de qualquer seed anterior).
- Duração em dias corridos é convertida em horas multiplicando por 8
  (aproximação combinada explicitamente, mesmo sabendo que "dias" no MS
  Project é prazo de calendário, não esforço).
- Atividades "resumo" (que têm subatividades) NÃO recebem horas previstas —
  só atividades-folha, pra não inflar relatórios que somam horas.
- A única tarefa com duração "0 hrs" vira um Marco, não uma Atividade.
- Toda atividade que o Excel mostra "em andamento" ganha um relato automático
  de importação antes de ter o status alterado (o Ergon PM exige isso).

ANTES DE RODAR
---------------
1. Rode a migration_006 no seu banco (Supabase ou local) se ainda não rodou —
   veja db/migration_006_relatos_e_requisito_tr_atividade.sql. Sem isso a
   importação falha logo no início (o script confere isso sozinho).
2. Instale as dependências:  pip install openpyxl requests
3. Tenha o backend do Ergon PM rodando e acessível (ex: docker compose up).

COMO RODAR
-----------
    python importar_cronograma_pcr.py --xlsx "ModeloCronograma.xlsx"

Por padrão aponta para http://localhost:8000/api. Se o backend estiver em
outro endereço:

    python importar_cronograma_pcr.py --xlsx "ModeloCronograma.xlsx" --api-url "http://SEU-HOST:8000/api"

O script pede uma confirmação explícita antes de gravar qualquer coisa
(pule com --yes se for rodar de forma não-interativa). Ele cria um projeto
NOVO — não mexe em nenhum projeto existente — então rodar duas vezes por
engano cria dois projetos duplicados; ele avisa e pede confirmação extra se
já existir um projeto com o mesmo nome/cliente.

Ao final, grava um relatório em JSON (import_result.json, na mesma pasta)
com tudo que foi criado, os avisos e eventuais erros.
"""
import argparse
import json
import re
import sys
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta a biblioteca 'openpyxl'. Instale com: pip install openpyxl")
try:
    import requests
except ImportError:
    sys.exit("Falta a biblioteca 'requests'. Instale com: pip install requests")


# ============================================================================
# Configuração de mapeamento (ajuste aqui se o WBS do seu cronograma mudar)
# ============================================================================
ETAPA_NOME = {
    "1.1": "Planejamento",
    "1.2": "Iniciação",
    "1.3": "Execução - Fase 1",
    "1.4": "Encerramento",
}

FRENTE_POR_EDT2 = {
    "1.3.1": "gestao", "1.3.2": "gestao",
    "1.3.3": "levantamento_processos",
    "1.3.4": "parametrizacao",
    "1.3.5": "migracao_dados",
    "1.3.6": "customizacoes",
    "1.3.7": "integracoes",
    "1.3.8": "folha_pagamento",
    "1.3.9": "portal_servidor",
    "1.3.10": "esocial",
    "1.3.11": "treinamentos",
    "1.3.12": "entrada_producao",
    "1.3.13": "operacao_assistida",
}
FRENTE_NOME = {
    "gestao": "Gestão",
    "levantamento_processos": "Levantamento de Processos e Casos de Uso",
    "parametrizacao": "Parametrização",
    "migracao_dados": "Migração de Dados",
    "customizacoes": "Customizações",
    "integracoes": "Integrações",
    "folha_pagamento": "Folha de Pagamento",
    "portal_servidor": "Portal do Gestor e Servidor",
    "esocial": "eSocial",
    "treinamentos": "Treinamentos",
    "entrada_producao": "Entrada em Produção",
    "operacao_assistida": "Operação Assistida",
}
CORES = ["#2f5f8a", "#7a4fa0", "#1f7a4d", "#a15c0d", "#0d8a99", "#ab2f2f",
         "#4a6b1f", "#8a5ca1", "#556b8a", "#8a6d3a", "#3a7a8a", "#5a5a5a"]

TIPO_KEYWORDS = [
    ("Homologação", ["homologa"]),
    ("Teste", ["teste", "testes da consultoria"]),
    ("Programação", ["programaç", "desenvolvimento"]),
    ("Levantamento", ["levantamento", "análise do levantamento"]),
    ("Especificação de Customização", ["especificaç"]),
    ("Parametrização", ["parametrizaç"]),
    ("Extração", ["extração"]),
    ("Carga", [" carga "]),
    ("Treinamento", ["treinamento"]),
    ("Reunião", ["reunião", "kick off", "kickoff"]),
    ("Entrega", ["entrega", "instalação", "go live"]),
    ("Comparação de Folha", ["comparação"]),
    ("Ajustes/Correções", ["ajuste", "correç"]),
]

STATUS_EXIGE_RELATO = {"Em andamento", "Bloqueada", "Cancelada"}
TIPO_DEPENDENCIA_CODE = {"": "FS", "TI": "FS", "II": "SS", "TT": "FF", "IT": "SF"}

PROJETO_NOME_PADRAO = "Implantação Sistema Ergon"
PROJETO_CLIENTE_PADRAO = "PCR"


# ============================================================================
# Cliente HTTP simples
# ============================================================================
class ApiClient:
    def __init__(self, base_url):
        self.base_url = base_url.rstrip("/")

    def call(self, method, path, **kwargs):
        url = self.base_url + path
        try:
            r = requests.request(method, url, timeout=30, **kwargs)
        except requests.exceptions.ConnectionError as e:
            raise SystemExit(
                f"\nNão consegui conectar em {url}.\n"
                f"Confirme que o backend do Ergon PM está rodando e acessível nesse endereço "
                f"(use --api-url para apontar pra outro). Detalhe técnico: {e}"
            )
        if not r.ok:
            raise RuntimeError(f"{method} {path} -> HTTP {r.status_code}: {r.text[:500]}")
        if r.status_code == 204:
            return None
        return r.json()


def preflight_check(api):
    """Confere que a migration_006 já foi aplicada, antes de criar qualquer coisa."""
    try:
        api.call("GET", "/atividades?projeto_id=00000000-0000-0000-0000-000000000000")
    except RuntimeError as e:
        if "requisito_tr_id" in str(e) or "atividade_relato" in str(e):
            raise SystemExit(
                "\nO banco ainda não tem a migration_006 aplicada (falta a coluna "
                "'requisito_tr_id' e/ou a tabela 'atividade_relato').\n"
                "Rode db/migration_006_relatos_e_requisito_tr_atividade.sql no SQL Editor "
                "do Supabase (ou no seu Postgres) antes de importar.\n\n"
                f"Detalhe técnico: {e}"
            )
        raise SystemExit(f"\nErro inesperado ao checar o backend: {e}")


# ============================================================================
# Parsing do Excel
# ============================================================================
def indent_of(nome):
    return len(nome) - len(nome.lstrip(" "))


def edt_dots(edt):
    return str(edt).count(".") if edt is not None else 0


def parse_date(s):
    if not s:
        return None
    m = re.search(r"(\d{2})/(\d{2})/(\d{2})$", str(s))
    if not m:
        return None
    d, mo, y = m.groups()
    return date(2000 + int(y), int(mo), int(d))


def parse_duracao_horas(s):
    """Retorna (horas, aproximado:bool) ou (None, False) se não reconhecido."""
    if not s:
        return None, False
    txt = str(s).strip().replace(",", ".")
    m = re.match(r"([\d.]+)\s*(hrs?|dias?|dia|sem(?:s|anas)?)$", txt, re.IGNORECASE)
    if not m:
        return None, False
    val = float(m.group(1))
    unidade = m.group(2).lower()
    if unidade.startswith("hr"):
        return val, False
    if unidade.startswith("dia"):
        return val * 8.0, True
    if unidade.startswith("sem"):
        return val * 5 * 8.0, True
    return None, False


def parse_lag_horas(txt):
    m = re.search(r"([+-]?\d+(?:[.,]\d+)?)\s*(hrs?|dias?|dia|sem(?:s|anas)?)", txt, re.IGNORECASE)
    if not m:
        return 0.0
    val = float(m.group(1).replace(",", "."))
    unidade = m.group(2).lower()
    if unidade.startswith("hr"):
        return val
    if unidade.startswith("dia"):
        return val * 8.0
    if unidade.startswith("sem"):
        return val * 5 * 8.0
    return 0.0


def parse_predecessoras(txt, warnings):
    out = []
    if not txt:
        return out
    for token in re.split(r"[;,]\s*", str(txt).strip()):
        if not token:
            continue
        m = re.match(r"(\d+)\s*(II|TT|IT|TI)?\s*([+-]\s*[\d.,]+\s*[a-zA-Zçã]+)?", token)
        if not m:
            warnings.append(f"predecessora não reconhecida: {token!r}")
            continue
        excel_id = int(m.group(1))
        tipo = TIPO_DEPENDENCIA_CODE.get(m.group(2) or "", "FS")
        lag = parse_lag_horas(m.group(3)) if m.group(3) else 0.0
        out.append({"excel_id": excel_id, "tipo": tipo, "lag_horas": lag})
    return out


def parse_recursos(txt, recurso_info):
    if not txt:
        return []
    out = []
    for token in str(txt).split(";"):
        token = token.strip()
        if not token:
            continue
        nome_limpo = re.sub(r"\[\d+%\]\s*$", "", token).strip()
        out.append(nome_limpo)
        if nome_limpo not in recurso_info:
            if nome_limpo.startswith("PCR"):
                recurso_info[nome_limpo] = {"nome": nome_limpo, "tipo_vinculo": "Cliente", "empresa": "PCR"}
            elif nome_limpo.startswith("Techne"):
                recurso_info[nome_limpo] = {"nome": nome_limpo, "tipo_vinculo": "Techne", "empresa": "Techne"}
            else:
                recurso_info[nome_limpo] = {"nome": nome_limpo, "tipo_vinculo": "Terceirizado", "empresa": None}
    return out


class Node:
    def __init__(self, row):
        self.edt = row[0]
        self.excel_id = row[1]
        nome_raw = row[2] or ""
        self.nome = nome_raw.strip()
        self.indent = indent_of(nome_raw)
        self.dots = edt_dots(self.edt)
        self.dtini = parse_date(row[3])
        self.dtfim = parse_date(row[4])
        self.duracao_txt = row[5]
        self.pct = row[6] or 0
        self.predecessoras_txt = row[9]
        self.recursos_txt = row[10]
        self.children = []
        self.parent = None
        self.is_summary = False
        self.etapa_key = None
        self.frente_key = None
        self.atividade_pai_node = None


def carregar_planilha(xlsx_path, sheet_name, warnings):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        # tenta achar a primeira aba com dado de verdade
        for name in wb.sheetnames:
            if wb[name].max_row > 1:
                sheet_name = name
                break
    ws = wb[sheet_name]
    raw_rows = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 13)]
        if all(v is None for v in vals):
            continue
        raw_rows.append(vals)

    nodes = [Node(row) for row in raw_rows]
    for i, node in enumerate(nodes):
        if i + 1 < len(nodes) and nodes[i + 1].indent > node.indent:
            node.is_summary = True

    stack = []
    for node in nodes:
        while stack and stack[-1].indent >= node.indent:
            stack.pop()
        if stack:
            node.parent = stack[-1]
            stack[-1].children.append(node)
        stack.append(node)

    if not nodes:
        raise SystemExit(f"Não encontrei nenhuma linha de dado na aba '{sheet_name}' de {xlsx_path}.")

    root = nodes[0]
    etapas_nodes = [n for n in nodes if n.dots == 1]

    for etapa_node in etapas_nodes:
        def walk(n, ek):
            n.etapa_key = ek
            for c in n.children:
                walk(c, ek)
        walk(etapa_node, etapa_node.edt)

    for node in nodes:
        if node.dots >= 2 and str(node.edt).startswith("1.3."):
            edt2 = ".".join(str(node.edt).split(".")[:3])
            fk = FRENTE_POR_EDT2.get(edt2, "gestao")

            def walk2(n, f):
                n.frente_key = f
                for c in n.children:
                    walk2(c, f)
            walk2(node, fk)
        elif node.dots >= 2 and str(node.edt).startswith(("1.2.", "1.4.")):
            node.frente_key = "gestao"

    for node in nodes:
        if node.dots < 2:
            continue
        anc = node.parent
        while anc is not None and anc.dots < 2:
            anc = anc.parent
        node.atividade_pai_node = anc

    def eh_marco(node):
        if node.is_summary:
            return False
        txt = str(node.duracao_txt or "").strip().lower()
        return bool(re.match(r"^0([.,]0+)?\s*hrs?$", txt))

    milestone_nodes = [n for n in nodes if n.dots >= 2 and eh_marco(n)]
    milestone_set = set(id(n) for n in milestone_nodes)
    atividade_nodes = [n for n in nodes if n.dots >= 2 and id(n) not in milestone_set]

    return root, etapas_nodes, atividade_nodes, milestone_nodes, nodes


# ============================================================================
# Carga via API
# ============================================================================
def confirmar_ou_sair(mensagem, auto_yes):
    if auto_yes:
        return
    resposta = input(f"{mensagem} Digite CONFIRMAR para prosseguir: ").strip()
    if resposta.upper() != "CONFIRMAR":
        sys.exit("Cancelado pelo usuário.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", required=True, help="Caminho do arquivo .xlsx exportado do MS Project.")
    parser.add_argument("--sheet", default="Plan1", help="Nome da aba com os dados (padrão: Plan1).")
    parser.add_argument("--api-url", default="http://localhost:8000/api", help="URL base da API do Ergon PM.")
    parser.add_argument("--projeto-nome", default=PROJETO_NOME_PADRAO)
    parser.add_argument("--projeto-cliente", default=PROJETO_CLIENTE_PADRAO)
    parser.add_argument("--yes", action="store_true", help="Não pede confirmação interativa (uso não-interativo).")
    parser.add_argument("--relatorio", default="import_result.json", help="Onde salvar o relatório final em JSON.")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        sys.exit(f"Arquivo não encontrado: {xlsx_path}")

    warnings = []
    api = ApiClient(args.api_url)

    print(f"Lendo {xlsx_path} (aba '{args.sheet}')...")
    root, etapas_nodes, atividade_nodes, milestone_nodes, all_nodes = carregar_planilha(xlsx_path, args.sheet, warnings)
    print(f"  Projeto (linha raiz): {root.nome!r} — {root.dtini} a {root.dtfim}")
    print(f"  Etapas encontradas: {[n.edt for n in etapas_nodes]}")
    print(f"  Atividades a criar: {len(atividade_nodes)}")
    print(f"  Marcos a criar (duração 0 hrs): {len(milestone_nodes)}")

    print(f"\nChecando backend em {args.api_url} ...")
    preflight_check(api)
    print("  OK — migration_006 detectada.")

    existentes = api.call("GET", "/projetos")
    ja_existe = [p for p in existentes if p.get("nome") == args.projeto_nome and p.get("cliente") == args.projeto_cliente]
    if ja_existe:
        confirmar_ou_sair(
            f"\nJá existe um projeto '{args.projeto_nome}' / cliente '{args.projeto_cliente}' "
            f"(id {ja_existe[0]['id']}). Rodar de novo cria um projeto DUPLICADO, não atualiza o existente.",
            args.yes,
        )
    else:
        confirmar_ou_sair(
            f"\nIsso vai criar um novo projeto '{args.projeto_nome}' (cliente '{args.projeto_cliente}') "
            f"com {len(atividade_nodes)} atividades, {len(etapas_nodes)} etapas, {len(FRENTE_NOME)} frentes "
            f"e {len(milestone_nodes)} marco(s) em {args.api_url}.",
            args.yes,
        )

    recurso_info = {}
    for node in atividade_nodes:
        parse_recursos(node.recursos_txt, recurso_info)

    print("\nCriando projeto...")
    projeto = api.call("POST", "/projetos", json={
        "nome": args.projeto_nome,
        "cliente": args.projeto_cliente,
        "sigla": args.projeto_cliente,
        "descricao": "Importado automaticamente do cronograma real (export MS Project em .xlsx).",
        "data_inicio": root.dtini.isoformat() if root.dtini else None,
        "data_fim_prevista": root.dtfim.isoformat() if root.dtfim else None,
        "horas_dia_util": 8,
    })
    projeto_id = projeto["id"]
    print("  Projeto criado:", projeto_id)

    print("Criando etapas...")
    etapa_id_by_key = {}
    for i, etapa_node in enumerate(etapas_nodes, start=1):
        descendentes = [n for n in all_nodes if n.etapa_key == etapa_node.edt and n.dtini]
        dtini = min((n.dtini for n in descendentes), default=etapa_node.dtini)
        dtfim = max((n.dtfim for n in descendentes if n.dtfim), default=etapa_node.dtfim)
        nome_etapa = ETAPA_NOME.get(etapa_node.edt, etapa_node.nome)
        etapa = api.call("POST", "/etapas", json={
            "projeto_id": projeto_id,
            "numero": i,
            "nome": nome_etapa,
            "descricao": f"Importado do cronograma real (EDT {etapa_node.edt} — \"{etapa_node.nome}\").",
            "data_inicio_prev": dtini.isoformat() if dtini else None,
            "data_fim_prev": dtfim.isoformat() if dtfim else None,
        })
        etapa_id_by_key[etapa_node.edt] = etapa["id"]
    print(f"  {len(etapa_id_by_key)} etapas criadas.")

    print("Criando frentes de trabalho...")
    frente_id_by_key = {}
    for i, (key, nome) in enumerate(FRENTE_NOME.items()):
        frente = api.call("POST", "/frentes", json={
            "projeto_id": projeto_id, "nome": nome, "cor_hex": CORES[i % len(CORES)], "ordem": i + 1,
        })
        frente_id_by_key[key] = frente["id"]
    print(f"  {len(frente_id_by_key)} frentes criadas.")

    print("Cadastrando recursos (responsáveis)...")
    recurso_id_by_nome = {r["nome"]: r["id"] for r in api.call("GET", "/recursos")}
    for nome, info in recurso_info.items():
        if nome in recurso_id_by_nome:
            continue
        rec = api.call("POST", "/recursos", json={
            "nome": info["nome"], "tipo_vinculo": info["tipo_vinculo"], "empresa": info["empresa"],
        })
        recurso_id_by_nome[nome] = rec["id"]
    print(f"  {len(recurso_id_by_nome)} recursos disponíveis (novos + já existentes).")

    print("Garantindo tipos de atividade elementar...")
    tipo_id_by_nome = {t["nome"]: t["id"] for t in api.call("GET", "/tipos-atividade")}
    for nome_tipo, _ in TIPO_KEYWORDS:
        if nome_tipo not in tipo_id_by_nome:
            t = api.call("POST", "/tipos-atividade", json={"nome": nome_tipo, "ordem": len(tipo_id_by_nome) + 1})
            tipo_id_by_nome[nome_tipo] = t["id"]

    def classificar_tipo(nome_tarefa):
        nome_lower = nome_tarefa.lower()
        for nome_tipo, kws in TIPO_KEYWORDS:
            for kw in kws:
                if kw in nome_lower:
                    return tipo_id_by_nome[nome_tipo]
        return None

    print(f"\nCriando {len(atividade_nodes)} atividades (pode levar alguns minutos)...")
    excel_id_to_atividade_id = {}
    hoje = datetime.now().date().isoformat()
    horas_aproximadas = 0
    criadas = 0
    erros = []

    for idx, node in enumerate(atividade_nodes, start=1):
        frente_key = node.frente_key or "gestao"
        frente_id = frente_id_by_key[frente_key]
        etapa_id = etapa_id_by_key[node.etapa_key]

        horas, aprox = parse_duracao_horas(node.duracao_txt)
        if node.is_summary:
            horas = None
        elif aprox:
            horas_aproximadas += 1

        recursos = parse_recursos(node.recursos_txt, recurso_info)
        techne = next((r for r in recursos if r.startswith("Techne")), None)
        cliente = next((r for r in recursos if r.startswith(args.projeto_cliente)), None)
        extras = [r for r in recursos if r not in (techne, cliente)]

        pct = round((node.pct or 0) * 100)
        if pct >= 100:
            status_final = "Concluída"
        elif pct <= 0:
            status_final = "Não iniciada"
        else:
            status_final = "Em andamento"

        pai_atividade_id = None
        if node.atividade_pai_node is not None:
            pai_atividade_id = excel_id_to_atividade_id.get(node.atividade_pai_node.excel_id)

        payload = {
            "projeto_id": projeto_id, "etapa_id": etapa_id, "frente_trabalho_id": frente_id,
            "atividade_pai_id": pai_atividade_id,
            "codigo_wbs": node.edt, "nome": node.nome[:250],
            "tipo_atividade_elementar_id": classificar_tipo(node.nome),
            "responsavel_techne_id": recurso_id_by_nome.get(techne) if techne else None,
            "responsavel_cliente_id": recurso_id_by_nome.get(cliente) if cliente else None,
            "prazo_horas": horas,
            "dtini_prev": node.dtini.isoformat() if node.dtini else None,
            "dtfim_prev": node.dtfim.isoformat() if node.dtfim else None,
            "percentual_concluido": pct,
            "observacoes": (
                f"Importado do cronograma MS Project (Excel), EDT {node.edt}, Id original {node.excel_id}."
                + (" Duração original em dias corridos, convertida em horas por aproximação (dias × 8h)." if aprox else "")
            ),
        }
        try:
            atividade = api.call("POST", "/atividades", json=payload)
        except RuntimeError as e:
            erros.append({"edt": node.edt, "nome": node.nome, "erro": str(e)})
            continue
        excel_id_to_atividade_id[node.excel_id] = atividade["id"]
        criadas += 1
        if criadas % 50 == 0:
            print(f"  ... {criadas}/{len(atividade_nodes)}")

        for extra_nome in extras:
            rid = recurso_id_by_nome.get(extra_nome)
            if rid:
                try:
                    api.call("POST", f"/atividades/{atividade['id']}/recursos", json={"recurso_id": rid})
                except RuntimeError as e:
                    erros.append({"edt": node.edt, "nome": "recurso extra", "erro": str(e)})

        if status_final in STATUS_EXIGE_RELATO:
            api.call("POST", f"/atividades/{atividade['id']}/relatos", json={
                "autor_nome": "Importação automática",
                "texto": f"Importado do cronograma MS Project (Excel) em {hoje}. "
                         f"Percentual concluído no arquivo original: {pct}%.",
            })
            api.call("PUT", f"/atividades/{atividade['id']}", json={"status": status_final})
        elif status_final == "Concluída":
            api.call("PUT", f"/atividades/{atividade['id']}", json={"status": status_final})

    print(f"\nAtividades criadas: {criadas} de {len(atividade_nodes)}")
    if erros:
        print(f"  ERROS: {len(erros)} (ver relatório final)")

    print("Criando dependências entre atividades...")
    dep_criadas = 0
    dep_erros = []
    for node in atividade_nodes:
        preds = parse_predecessoras(node.predecessoras_txt, warnings)
        if not preds:
            continue
        ativ_id = excel_id_to_atividade_id.get(node.excel_id)
        if not ativ_id:
            continue
        for p in preds:
            pred_ativ_id = excel_id_to_atividade_id.get(p["excel_id"])
            if not pred_ativ_id:
                warnings.append(
                    f"dependência órfã: atividade Id {node.excel_id} -> predecessora Id {p['excel_id']} (não encontrada)"
                )
                continue
            try:
                api.call("POST", f"/atividades/{ativ_id}/dependencias", json={
                    "predecessora_id": pred_ativ_id, "tipo": p["tipo"], "lag_horas": p["lag_horas"],
                })
                dep_criadas += 1
            except RuntimeError as e:
                dep_erros.append({"edt": node.edt, "erro": str(e)})
    print(f"  {dep_criadas} dependências criadas.")

    print("Extraindo marcos...")
    marcos_criados = 0
    for node in milestone_nodes:
        data_prevista = node.dtfim or node.dtini
        if not data_prevista:
            warnings.append(f"marco '{node.nome}' sem data — não importado.")
            continue
        pct = round((node.pct or 0) * 100)
        api.call("POST", "/marcos", json={
            "projeto_id": projeto_id,
            "etapa_id": etapa_id_by_key.get(node.etapa_key),
            "nome": node.nome[:250],
            "descricao": f"Importado do cronograma MS Project (Excel), EDT {node.edt}, Id original {node.excel_id}.",
            "data_prevista": data_prevista.isoformat(),
            "data_real": data_prevista.isoformat() if pct >= 100 else None,
        })
        marcos_criados += 1
    print(f"  {marcos_criados} marcos criados.")

    if warnings:
        print(f"\nAvisos ({len(warnings)}):")
        for w in warnings[:30]:
            print("  -", w)

    summary = {
        "projeto_id": projeto_id,
        "atividades_criadas": criadas,
        "atividades_total": len(atividade_nodes),
        "dependencias_criadas": dep_criadas,
        "recursos": len(recurso_id_by_nome),
        "horas_aproximadas": horas_aproximadas,
        "marcos_criados": marcos_criados,
        "erros_atividades": erros,
        "erros_dependencias": dep_erros,
        "avisos": warnings,
    }
    Path(args.relatorio).write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"\nRelatório salvo em {args.relatorio}")
    print("\nOK — importação concluída. Abra o Ergon PM e confira o projeto criado.")


if __name__ == "__main__":
    main()
