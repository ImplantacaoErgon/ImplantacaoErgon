"""
Relatório de Horas dos Recursos (apontamento de horas em "Minhas
atividades") — gerado a partir do Dashboard, ao lado do Relatório
Executivo (IA).

Diferente do Relatório Executivo, este é 100% determinístico (só monta e
formata dados que já existem em `atividade_horas_dia` — nenhuma IA
envolvida) e gerado na hora, sem persistir nada no banco: os filtros
(recurso, período) mudam a cada geração, então não faz sentido guardar
histórico como no relatório executivo.

Escopo: como "Minhas atividades" (ver docstring de
`backend/app/minhas_atividades.py`), este relatório traz o apontamento em
TODOS os projetos, não só o selecionado no cabeçalho — é uma reconciliação
de horas trabalhadas na empresa como um todo, por isso cada atividade traz
consigo o projeto a que pertence. Só recursos com `controla_horas = true`
registram apontamento em "Minhas atividades" — na prática só eles aparecem
neste relatório, mas a consulta em si não filtra por essa marca (ela é
aplicada no combo de filtro do frontend; um recurso que teve a marca
desativada depois de já ter apontamentos continua aparecendo aqui, o que é
o comportamento correto para não perder histórico).

`coletar_dados()` faz a consulta e agrupa em memória por recurso →
atividade → dias, com subtotais em cada nível; `gerar_pdf_bytes()` e
`gerar_planilha_bytes()` formatam esse mesmo dicionário em PDF (reportlab)
ou XLSX (openpyxl) — dois formatos, um dado só.
"""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

from . import db

COR_CABECALHO = colors.HexColor("#1a3d5c")
COR_CABECALHO_HEX = "1A3D5C"
COR_FAIXA = colors.HexColor("#f2f5f8")


def _fmt_data_br(d):
    if d is None:
        return ""
    if isinstance(d, str):
        try:
            d = date.fromisoformat(d)
        except ValueError:
            return d
    return d.strftime("%d/%m/%Y")


def _fmt_horas(v):
    return f"{float(v):.1f}".replace(".", ",") if v is not None else "—"


def coletar_dados(recurso_id, data_ini, data_fim):
    """Monta a estrutura completa do relatório: lista de recursos, cada um
    com suas atividades (com subtotal previsto/realizado) e, dentro de cada
    atividade, os lançamentos diários. `recurso_id` None/vazio = todos os
    recursos que tiverem ao menos um lançamento no período."""
    filtro_recurso = f"AND hd.recurso_id = {db.q(recurso_id)}" if recurso_id else ""
    linhas = db.fetch_all(f"""
        SELECT hd.data, hd.horas_previstas, hd.horas_realizadas, hd.confirmado,
               r.id AS recurso_id, r.nome AS recurso_nome, r.tipo_vinculo,
               a.id AS atividade_id, a.nome AS atividade_nome, a.codigo_wbs,
               p.sigla AS projeto_sigla, p.nome AS projeto_nome,
               e.numero AS etapa_numero, e.nome AS etapa_nome, fr.nome AS frente_nome
        FROM atividade_horas_dia hd
        JOIN recursos r ON r.id = hd.recurso_id
        JOIN atividades a ON a.id = hd.atividade_id
        JOIN projetos p ON p.id = a.projeto_id
        JOIN etapas e ON e.id = a.etapa_id
        JOIN frentes_trabalho fr ON fr.id = a.frente_trabalho_id
        WHERE hd.data BETWEEN {db.q(data_ini)} AND {db.q(data_fim)}
        {filtro_recurso}
        ORDER BY r.nome, p.sigla NULLS LAST, a.codigo_wbs NULLS LAST, a.nome, hd.data
    """)

    consultores = {}
    for l in linhas:
        rid = l["recurso_id"]
        c = consultores.setdefault(rid, {
            "recurso_id": rid, "recurso_nome": l["recurso_nome"], "tipo_vinculo": l["tipo_vinculo"],
            "atividades": {}, "total_previsto": 0.0, "total_realizado": 0.0,
        })
        aid = l["atividade_id"]
        at = c["atividades"].setdefault(aid, {
            "atividade_id": aid, "atividade_nome": l["atividade_nome"], "codigo_wbs": l["codigo_wbs"],
            "projeto_sigla": l["projeto_sigla"], "projeto_nome": l["projeto_nome"],
            "etapa_numero": l["etapa_numero"], "etapa_nome": l["etapa_nome"], "frente_nome": l["frente_nome"],
            "dias": [], "total_previsto": 0.0, "total_realizado": 0.0,
        })
        previsto = float(l["horas_previstas"]) if l["horas_previstas"] is not None else 0.0
        realizado = float(l["horas_realizadas"]) if (l["confirmado"] and l["horas_realizadas"] is not None) else 0.0
        at["dias"].append({
            "data": l["data"], "horas_previstas": l["horas_previstas"],
            "horas_realizadas": l["horas_realizadas"] if l["confirmado"] else None,
            "confirmado": bool(l["confirmado"]),
        })
        at["total_previsto"] += previsto
        at["total_realizado"] += realizado
        c["total_previsto"] += previsto
        c["total_realizado"] += realizado

    resultado = []
    for c in consultores.values():
        c["atividades"] = sorted(
            c["atividades"].values(),
            key=lambda x: (x["projeto_sigla"] or x["projeto_nome"] or "", x["codigo_wbs"] or "", x["atividade_nome"])
        )
        resultado.append(c)
    resultado.sort(key=lambda x: x["recurso_nome"])

    return {
        "periodo": {"inicio": data_ini, "fim": data_fim},
        "recurso_filtrado": recurso_id or None,
        "consultores": resultado,
        "total_geral_previsto": sum(c["total_previsto"] for c in resultado),
        "total_geral_realizado": sum(c["total_realizado"] for c in resultado),
    }


# --------------------------------------------------------------------- PDF

def _estilos_pdf():
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("titulo", parent=base["Title"], fontSize=17, spaceAfter=4),
        "subtitulo": ParagraphStyle("subtitulo", parent=base["Normal"], fontSize=10,
                                     textColor=colors.HexColor("#555555"), alignment=TA_CENTER, spaceAfter=2),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontSize=13, spaceBefore=16, spaceAfter=6,
                              textColor=COR_CABECALHO),
        "h3": ParagraphStyle("h3", parent=base["Heading3"], fontSize=10.5, spaceBefore=8, spaceAfter=4,
                              textColor=colors.HexColor("#333333")),
        "corpo": ParagraphStyle("corpo", parent=base["Normal"], fontSize=9, leading=13),
        "rodape": ParagraphStyle("rodape", parent=base["Normal"], fontSize=7.5,
                                  textColor=colors.HexColor("#888888")),
    }


def _tabela_resumo_atividades(consultor, estilos):
    cabecalho = ["Atividade", "Projeto", "Etapa/Frente", "Previsto (h)", "Realizado (h)"]
    linhas = [cabecalho]
    for at in consultor["atividades"]:
        subtitulo = " / ".join(filter(None, [at.get("etapa_nome"), at.get("frente_nome")]))
        nome = (f"{at['codigo_wbs']} — " if at.get("codigo_wbs") else "") + at["atividade_nome"]
        linhas.append([
            Paragraph(nome, estilos["corpo"]),
            Paragraph(at.get("projeto_sigla") or at.get("projeto_nome") or "—", estilos["corpo"]),
            Paragraph(subtitulo or "—", estilos["corpo"]),
            _fmt_horas(at["total_previsto"]), _fmt_horas(at["total_realizado"]),
        ])
    linhas.append(["Total do recurso", "", "", _fmt_horas(consultor["total_previsto"]),
                    _fmt_horas(consultor["total_realizado"])])
    t = Table(linhas, colWidths=[6.5 * cm, 2.3 * cm, 3.8 * cm, 2.2 * cm, 2.2 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), COR_CABECALHO),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (3, 0), (4, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dddddd")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, COR_FAIXA]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e8edf2")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _tabela_detalhe_dias(consultor, estilos):
    cabecalho = ["Data", "Atividade", "Projeto", "Previsto (h)", "Realizado (h)", "Confirmado"]
    linhas = [cabecalho]
    for at in consultor["atividades"]:
        nome = (f"{at['codigo_wbs']} — " if at.get("codigo_wbs") else "") + at["atividade_nome"]
        for d in at["dias"]:
            linhas.append([
                _fmt_data_br(d["data"]), Paragraph(nome, estilos["corpo"]),
                at.get("projeto_sigla") or at.get("projeto_nome") or "—",
                _fmt_horas(d["horas_previstas"]),
                _fmt_horas(d["horas_realizadas"]) if d["confirmado"] else "—",
                "Sim" if d["confirmado"] else "Não",
            ])
    t = Table(linhas, colWidths=[2.1 * cm, 6.3 * cm, 2.3 * cm, 2.1 * cm, 2.1 * cm, 2.1 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5c7a90")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (3, 0), (5, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dddddd")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, COR_FAIXA]),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def gerar_pdf_bytes(dados):
    estilos = _estilos_pdf()
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        topMargin=1.8 * cm, bottomMargin=1.6 * cm, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
        title="Relatório de Horas dos Recursos",
    )
    periodo_txt = f"{_fmt_data_br(dados['periodo']['inicio'])} a {_fmt_data_br(dados['periodo']['fim'])}"
    consultores = dados["consultores"]
    filtro_txt = consultores[0]["recurso_nome"] if (dados["recurso_filtrado"] and len(consultores) == 1) else "Todos os recursos"

    flow = [
        Paragraph("Relatório de Horas dos Recursos", estilos["titulo"]),
        Paragraph(f"Período: {periodo_txt} · Recurso(s): {filtro_txt}", estilos["subtitulo"]),
        Paragraph("Apontamento de horas (previsto × realizado) lançado em \"Minhas atividades\", em todos os projetos.",
                  estilos["subtitulo"]),
        Spacer(1, 6),
        HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#cccccc")),
    ]

    if not consultores:
        flow.append(Spacer(1, 20))
        flow.append(Paragraph("Nenhum lançamento de horas encontrado para os filtros selecionados.", estilos["corpo"]))
    else:
        for c in consultores:
            flow.append(Paragraph(f"{c['recurso_nome']}  ({c['tipo_vinculo'] or '—'})", estilos["h2"]))
            flow.append(Paragraph("Resumo por atividade", estilos["h3"]))
            flow.append(_tabela_resumo_atividades(c, estilos))
            flow.append(Spacer(1, 8))
            flow.append(Paragraph("Detalhe diário", estilos["h3"]))
            flow.append(_tabela_detalhe_dias(c, estilos))
            flow.append(Spacer(1, 14))
        flow.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#cccccc")))
        flow.append(Spacer(1, 6))
        total_linha = Table(
            [["Total geral (todos os recursos listados)", _fmt_horas(dados["total_geral_previsto"]),
              _fmt_horas(dados["total_geral_realizado"])]],
            colWidths=[10.6 * cm, 2.2 * cm, 2.2 * cm],
        )
        total_linha.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#e8edf2")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9.5),
            ("ALIGN", (1, 0), (2, 0), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        flow.append(total_linha)

    flow.append(Spacer(1, 14))
    flow.append(Paragraph(
        "\"Realizado\" só considera dias já confirmados pelo recurso em \"Minhas atividades\". "
        "Estes números são independentes das horas previstas/realizadas registradas na atividade do Cronograma "
        "(que são um total único por atividade, não dividido entre os participantes).",
        estilos["rodape"],
    ))
    doc.build(flow)
    return buf.getvalue()


# ---------------------------------------------------------------- Planilha

def _estilo_cabecalho(ws, linha, n_colunas, cor_hex=COR_CABECALHO_HEX):
    for col in range(1, n_colunas + 1):
        cel = ws.cell(row=linha, column=col)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill(start_color=cor_hex, end_color=cor_hex, fill_type="solid")
        cel.alignment = Alignment(vertical="center")


def _autofit(ws, larguras):
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def gerar_planilha_bytes(dados):
    wb = Workbook()

    # ---- aba Resumo: uma linha por atividade por consultor, com subtotal ----
    ws1 = wb.active
    ws1.title = "Resumo"
    cabecalho = ["Recurso", "Vínculo", "Atividade", "Código", "Projeto", "Etapa", "Frente",
                 "Previsto (h)", "Realizado (h)"]
    ws1.append(cabecalho)
    _estilo_cabecalho(ws1, 1, len(cabecalho))
    linha = 2
    for c in dados["consultores"]:
        for at in c["atividades"]:
            ws1.append([
                c["recurso_nome"], c["tipo_vinculo"] or "", at["atividade_nome"], at.get("codigo_wbs") or "",
                at.get("projeto_sigla") or at.get("projeto_nome") or "", at.get("etapa_nome") or "",
                at.get("frente_nome") or "", round(at["total_previsto"], 2), round(at["total_realizado"], 2),
            ])
            linha += 1
        ws1.append([f"Total — {c['recurso_nome']}", "", "", "", "", "", "",
                    round(c["total_previsto"], 2), round(c["total_realizado"], 2)])
        for col in range(1, len(cabecalho) + 1):
            ws1.cell(row=linha, column=col).font = Font(bold=True)
        linha += 1
    ws1.append(["Total geral", "", "", "", "", "", "",
                round(dados["total_geral_previsto"], 2), round(dados["total_geral_realizado"], 2)])
    for col in range(1, len(cabecalho) + 1):
        ws1.cell(row=linha, column=col).font = Font(bold=True)
        ws1.cell(row=linha, column=col).fill = PatternFill(start_color="E8EDF2", end_color="E8EDF2", fill_type="solid")
    _autofit(ws1, [24, 12, 34, 10, 12, 20, 18, 13, 14])
    ws1.freeze_panes = "A2"

    # ---- aba Detalhe: um lançamento diário por linha ----
    ws2 = wb.create_sheet("Detalhe")
    cabecalho2 = ["Recurso", "Vínculo", "Data", "Atividade", "Código", "Projeto",
                  "Previsto (h)", "Realizado (h)", "Confirmado"]
    ws2.append(cabecalho2)
    _estilo_cabecalho(ws2, 1, len(cabecalho2), cor_hex="5C7A90")
    for c in dados["consultores"]:
        for at in c["atividades"]:
            for d in at["dias"]:
                ws2.append([
                    c["recurso_nome"], c["tipo_vinculo"] or "", _fmt_data_br(d["data"]), at["atividade_nome"],
                    at.get("codigo_wbs") or "", at.get("projeto_sigla") or at.get("projeto_nome") or "",
                    round(float(d["horas_previstas"]), 2) if d["horas_previstas"] is not None else None,
                    round(float(d["horas_realizadas"]), 2) if d["confirmado"] and d["horas_realizadas"] is not None else None,
                    "Sim" if d["confirmado"] else "Não",
                ])
    _autofit(ws2, [24, 12, 12, 34, 10, 12, 13, 14, 12])
    ws2.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
