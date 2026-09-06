"""
Extração heurística de requisitos a partir de um Termo de Referência (TR).

Não usa nenhum modelo de IA em tempo de execução (o backend não tem acesso
a nenhuma API de LLM) — é um parser baseado em regras sobre a estrutura
típica dos TRs de licitação pública brasileiros para sistemas de software:

  1. Um grande catálogo de "funcionalidades", organizado em módulos
     numerados com título em CAIXA ALTA ("1. ABRANGÊNCIA...",
     "2. FUNCIONALIDADES GERAIS...", "4. GESTÃO DA FOLHA...", etc.),
     cada um contendo itens numerados hierarquicamente (2.2.1, 2.3.4.2...)
     — tratados aqui como requisitos **Funcionais**.
  2. Um anexo separado, geralmente chamado "Requisitos Técnicos da
     Solução" / "Requisitos Não Funcionais" / "Requisitos de Desempenho",
     organizado por categoria (uso, armazenamento, desempenho,
     segurança, integração, disponibilidade...) — tratado aqui como
     requisitos **Não Funcionais**.

O algoritmo:
  - Converte o arquivo (HTML, TXT, DOCX ou PDF) em texto simples, uma
    linha por parágrafo/célula.
  - Localiza a sequência mais longa de títulos "N. TÍTULO EM CAIXA ALTA"
    numerados consecutivamente (1, 2, 3...) — essa é a zona funcional.
  - Localiza um título de anexo (linha isolada no formato
    "ANEXO ... – <título>") cujo título bata com um vocabulário de
    requisitos técnicos/não funcionais — essa é a zona não funcional,
    limitada pelo próximo título de anexo.
  - Dentro de cada zona, extrai linhas numeradas hierarquicamente
    (N.N, N.N.N, N.N.N.N...). Uma linha curta (até 12 palavras), sem
    verbo de obrigação ("deve", "permitir", "possuir"...) e que tem
    itens-filhos logo depois, é tratada como um SUBTÍTULO (vira contexto
    de "módulo de origem" dos itens abaixo dela) em vez de um requisito
    em si — evita importar linhas como "2.2. Localização e busca de
    informações" como se fossem um requisito próprio.

É uma heurística, não uma leitura garantida — sempre precisa de revisão
humana antes de confirmar a importação (por isso a rota correspondente
em main.py devolve uma PRÉVIA, não insere direto no banco).
"""
import io
import re
import unicodedata


OBLIGATION_VERBS = re.compile(
    r"\b(deve|dever[aá]|devem|deverao|deverão|permitir|possibilitar|possuir|garantir|"
    r"assegurar|manter|gerar|emitir|disponibilizar|registrar|realizar|controlar|"
    r"monitorar|integrar|exportar|importar|compat[ií]vel|suportar|apresentar|calcular|"
    r"efetuar|proceder|viabilizar|proporcionar|bloquear|notificar|encaminhar)\b",
    re.IGNORECASE,
)

NF_KEYWORDS = [
    "REQUISITOS TECNICOS", "REQUISITOS TÉCNICOS",
    "REQUISITOS NAO FUNCIONAIS", "REQUISITOS NÃO FUNCIONAIS", "REQUISITOS NAO-FUNCIONAIS",
    "REQUISITOS DE DESEMPENHO", "REQUISITOS DE INFRAESTRUTURA",
]

TOP_HEADING_RE = re.compile(r"^(\d{1,2})\.\s*(.{4,90}?):?\s*$")
LEAF_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,3}){1,4})\.?\s+(.{2,}?)\s*$")
ANEXO_RE = re.compile(r"^ANEXO\s+\S{1,15}\s*[\-–—]\s*(.{3,120})$", re.IGNORECASE)


class TrParseError(Exception):
    pass


def _strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _norm(s):
    return _strip_accents(s).upper()


def _is_upper_heading(s):
    letters = [c for c in s if c.isalpha()]
    if len(letters) < 5:
        return False
    upper = sum(1 for c in letters if c.isupper())
    return upper / len(letters) > 0.85


# ---------------------------------------------------------------- extração de texto

def extract_text(raw_bytes, filename):
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext in ("html", "htm"):
        return _text_from_html(raw_bytes)
    if ext == "txt":
        return _decode(raw_bytes)
    if ext == "docx":
        return _text_from_docx(raw_bytes)
    if ext == "pdf":
        return _text_from_pdf(raw_bytes)
    raise TrParseError(
        f"Formato '.{ext}' não suportado para importação automática. "
        "Aceita: .html, .htm, .txt, .docx, .pdf."
    )


def _decode(raw_bytes):
    try:
        return raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw_bytes.decode("latin-1")


def _text_from_html(raw_bytes):
    from html.parser import HTMLParser

    class _Extractor(HTMLParser):
        def __init__(self):
            super().__init__()
            self.parts = []
            self.skip = 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style"):
                self.skip += 1
            if tag in ("br", "p", "div", "tr", "li", "h1", "h2", "h3", "h4", "h5", "table"):
                self.parts.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style"):
                self.skip = max(0, self.skip - 1)
            if tag in ("p", "div", "tr", "li", "h1", "h2", "h3", "h4", "h5", "table"):
                self.parts.append("\n")

        def handle_data(self, data):
            if not self.skip:
                self.parts.append(data)

    html = _decode(raw_bytes)
    p = _Extractor()
    p.feed(html)
    text = "".join(p.parts)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text


def _text_from_docx(raw_bytes):
    import docx

    doc = docx.Document(io.BytesIO(raw_bytes))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    return "\n".join(lines)


def _text_from_pdf(raw_bytes):
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


# ---------------------------------------------------------------- localização das zonas

def _find_functional_zone(lines):
    """Acha a maior sequência de títulos 'N. TÍTULO EM CAIXA ALTA' numerados
    consecutivamente (1,2,3...) — essa é o catálogo de funcionalidades."""
    candidates = []
    for i, line in enumerate(lines):
        s = line.strip()
        m = TOP_HEADING_RE.match(s)
        if m and _is_upper_heading(m.group(2)):
            candidates.append((i, int(m.group(1)), m.group(2).strip()))

    best, cur = [], []
    for c in candidates:
        if not cur:
            cur = [c] if c[1] in (1, 2) else []
        elif c[1] == cur[-1][1] + 1:
            cur.append(c)
        else:
            cur = [c] if c[1] in (1, 2) else []
        if len(cur) > len(best):
            best = cur
    return best  # lista de (linha, numero, titulo)


def _find_anexo_headings(lines):
    """Todas as linhas isoladas no formato 'ANEXO <id> – <título>' —
    reaproveitado tanto para achar o anexo de requisitos técnicos/não
    funcionais quanto como limite de segurança da zona funcional (evita
    que ela "vaze" para dentro de um anexo qualquer não reconhecido)."""
    heads = []
    for i, line in enumerate(lines):
        s = line.strip()
        if len(s) > 160:
            continue
        if ANEXO_RE.match(s):
            heads.append((i, s))
    return heads


def _find_non_functional_zone(anexo_heads, total_lines):
    """Acha, entre os títulos de anexo já localizados, um que bata com
    vocabulário de requisitos técnicos/não funcionais, limitado pelo
    próximo título de anexo (ou pelo fim do documento)."""
    nf_start = None
    for i, s in anexo_heads:
        if any(kw in _norm(s) for kw in NF_KEYWORDS):
            nf_start = i
            break
    if nf_start is None:
        return None

    nf_end = total_lines
    for i, _ in anexo_heads:
        if i > nf_start:
            nf_end = i
            break
    return (nf_start, nf_end)


# ---------------------------------------------------------------- extração dos itens

def _extract_leaf_items(lines, start_line, end_line):
    raw = []
    current_top = None
    for i in range(start_line, end_line):
        s = lines[i].strip()
        if not s:
            continue
        m = LEAF_RE.match(s)
        if not m:
            # pode ser o próprio título do módulo/categoria de 1 nível ("1. Requisitos de uso")
            m_top = TOP_HEADING_RE.match(s)
            if m_top:
                current_top = f"{m_top.group(1)}. {m_top.group(2).strip()}"
            continue
        codigo, texto = m.group(1), m.group(2).strip()
        depth = codigo.count(".") + 1
        if depth == 1:
            current_top = f"{codigo}. {texto}"
            continue
        raw.append({"codigo": codigo, "texto": texto, "depth": depth, "top": current_top})

    codes = {r["codigo"] for r in raw}

    def has_child(codigo):
        prefix = codigo + "."
        return any(c.startswith(prefix) for c in codes)

    final = []
    heading_stack = {}
    for r in raw:
        wc = len(r["texto"].split())
        is_label = has_child(r["codigo"]) and wc <= 12 and not OBLIGATION_VERBS.search(r["texto"])
        if is_label:
            heading_stack[r["depth"]] = f"{r['codigo']} {r['texto']}"
            for d in list(heading_stack):
                if d > r["depth"]:
                    del heading_stack[d]
            continue
        path = [r["top"]] if r["top"] else []
        for d in sorted(heading_stack):
            if d < r["depth"]:
                path.append(heading_stack[d])
        final.append({
            "codigo": r["codigo"],
            "texto": r["texto"],
            "modulo_origem": " > ".join(path) if path else None,
        })
    return final


def _titulo_from_texto(texto):
    m = re.match(r"^(.{20,140}?[\.\!\?])(\s|$)", texto)
    base = m.group(1) if m else texto[:140]
    base = base.strip()
    if len(base) < len(texto.strip()) and not base.endswith((".", "!", "?")):
        base += "…"
    return base


def _build_item(codigo, texto, modulo_origem, prefixo, tipo):
    return {
        "codigo": f"{prefixo}-{codigo}",
        "titulo": _titulo_from_texto(texto),
        "descricao": texto,
        "modulo_origem": modulo_origem,
        "tipo_requisito": tipo,
    }


# ---------------------------------------------------------------- API principal

def parse_tr_document(raw_bytes, filename):
    """Retorna {"itens": [...], "avisos": [...]}. Cada item já vem no
    formato pronto para inserir em requisitos_tr (falta só projeto_id,
    frente_trabalho_id, classificacao/atendimento/status default)."""
    text = extract_text(raw_bytes, filename)
    lines = text.split("\n")
    avisos = []

    anexo_heads = _find_anexo_headings(lines)
    func_zone = _find_functional_zone(lines)
    nf_zone = _find_non_functional_zone(anexo_heads, len(lines))

    itens = []

    if func_zone:
        start_line = func_zone[0][0]
        # limite da zona funcional: o anexo de requisitos não funcionais, se achado;
        # senão, o próximo título de anexo qualquer (evita "vazar" para dentro de um
        # anexo não reconhecido); senão, uma janela de segurança.
        proximo_anexo = next((i for i, _ in anexo_heads if i > func_zone[-1][0]), None)
        if nf_zone:
            end_line = nf_zone[0]
        elif proximo_anexo is not None:
            end_line = proximo_anexo
        else:
            end_line = func_zone[-1][0] + 400
        end_line = min(end_line, len(lines))
        for it in _extract_leaf_items(lines, start_line, end_line):
            itens.append(_build_item(it["codigo"], it["texto"], it["modulo_origem"], "RF", "Funcional"))
    else:
        avisos.append(
            "Não foi possível identificar automaticamente o catálogo de funcionalidades "
            "(nenhuma sequência de módulos numerados 1, 2, 3... em caixa alta foi encontrada)."
        )

    if nf_zone:
        start_line, end_line = nf_zone
        for it in _extract_leaf_items(lines, start_line, end_line):
            itens.append(_build_item(it["codigo"], it["texto"], it["modulo_origem"], "RNF", "Não Funcional"))
    else:
        avisos.append(
            "Não foi encontrado um anexo de 'Requisitos Técnicos'/'Requisitos Não Funcionais' — "
            "se o seu TR tiver essa seção com outro nome, os itens dela não foram capturados."
        )

    if not itens:
        raise TrParseError(
            "Não foi possível reconhecer a estrutura de requisitos deste documento. "
            "Ele precisa ter um catálogo de módulos numerados (ex: '4. GESTÃO DA FOLHA DE PAGAMENTO') "
            "com itens numerados hierarquicamente (ex: '4.3.2.'). Considere usar a importação por CSV."
        )

    # checagem de qualidade: numa lista real de requisitos funcionais, a maioria das
    # frases usa um verbo de obrigação/capacidade ("o sistema deve/deverá/permitir/
    # possuir..."). Se a proporção for baixa, é sinal de que o parser pode ter pego a
    # sequência numerada errada (ex: cláusulas de um Edital, não o TR de fato) — não
    # bloqueia a prévia, mas avisa com destaque para o usuário revisar com atenção.
    funcionais = [it for it in itens if it["tipo_requisito"] == "Funcional"]
    baixa_confianca = False
    if funcionais:
        com_verbo = sum(1 for it in funcionais if OBLIGATION_VERBS.search(it["descricao"]))
        if com_verbo / len(funcionais) < 0.35:
            baixa_confianca = True
            avisos.insert(0,
                "⚠ Só uma pequena parte dos itens encontrados parece descrever uma "
                "funcionalidade do sistema (frases como 'o sistema deve/deverá/permitir...'). "
                "Isso é um sinal de que este pode não ser o documento certo (ex: o Edital em vez "
                "do Anexo/Termo de Referência com as especificações técnicas), ou de que a "
                "estrutura não é a esperada. Revise a lista com atenção antes de confirmar."
            )

    # códigos duplicados dentro do próprio arquivo (raro, mas possível se o TR reaproveita numeração)
    vistos = {}
    for it in itens:
        vistos[it["codigo"]] = vistos.get(it["codigo"], 0) + 1
    repetidos = [c for c, n in vistos.items() if n > 1]
    if repetidos:
        avisos.append(
            f"{len(repetidos)} código(s) apareceram mais de uma vez no documento "
            f"(ex: {', '.join(repetidos[:5])}) — ao confirmar, o último sobrescreve os anteriores."
        )

    return {"itens": itens, "avisos": avisos, "baixa_confianca": baixa_confianca}
